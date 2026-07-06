from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from zijinhua_tekla.classifiers.box_main_material_segments import classify_box_main_material_segment_groups
from zijinhua_tekla.classifiers.box_part_spatial_relations import classify_box_part_spatial_relations
from zijinhua_tekla.classifiers.composite_main_material_segments import (
    CompositeSegmentType,
    classify_composite_main_material_segments,
)
from zijinhua_tekla.spatial_features import classify_appendage_clusters_from_bundle


MEMBER_SHEET_INDEX = 1
PART_SHEET_INDEX = 2

MEMBER_ID_COL = 1
MEMBER_TEAM_COL = 7
MEMBER_LINE_COL = 8

PART_MEMBER_COL = 1
PART_POSITION_COL = 2
PART_DELIVERY_TARGET_COL = 12


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def load_member_routes(ws) -> dict[str, dict[str, str]]:
    routes: dict[str, dict[str, str]] = {}
    for row in range(1, ws.max_row + 1):
        member = text(ws.cell(row, MEMBER_ID_COL).value)
        if not member or member in {"构件编号", "构件清单"}:
            continue
        team = text(ws.cell(row, MEMBER_TEAM_COL).value)
        line = text(ws.cell(row, MEMBER_LINE_COL).value)
        if team or line:
            routes[member] = {"team": team, "line": line}
    return routes


def load_member_snapshot(cache_dir: Path, member: str) -> dict[str, Any] | None:
    member_path = cache_dir / "members" / f"member_{member}.json"
    if not member_path.exists():
        return None
    return json.loads(member_path.read_text(encoding="utf-8"))


def algorithm_body_type(member_snapshot: dict[str, Any] | None) -> str:
    if not member_snapshot:
        return "UNKNOWN"
    classification = member_snapshot.get("Classification") or {}
    main_class = text(classification.get("MainClass")).upper()
    if main_class in {"1", "H", "BH", "H_BEAM"}:
        return "H"
    if main_class in {"2", "BOX"}:
        return "BOX"
    key_dimensions = text(classification.get("KeyDimensionsDisplay")).upper()
    if key_dimensions.startswith(("BH", "H")) and not key_dimensions.startswith("HP"):
        return "H"
    if key_dimensions.startswith("BOX"):
        return "BOX"
    return main_class or "UNKNOWN"


def load_algorithm_annotations(
    cache_dir: Path,
) -> tuple[set[str], dict[str, str], dict[tuple[str, str], dict[str, Any]]]:
    bundle = cache_dir / "tekla-body-bracket-export.bundle.json"
    if not bundle.exists():
        return set(), {}, {}

    data = json.loads(bundle.read_text(encoding="utf-8"))
    exported_members: set[str] = set()
    member_body_types: dict[str, str] = {}
    annotations: dict[tuple[str, str], dict[str, Any]] = {}

    for assembly in data.get("assemblies", []):
        metadata = assembly.get("metadata") or {}
        member = text(metadata.get("assemblyPosition"))
        if not member:
            continue
        exported_members.add(member)
        member_snapshot = load_member_snapshot(cache_dir, member)
        body_type = algorithm_body_type(member_snapshot)
        member_body_types[member] = body_type
        main_groups = classify_box_main_material_segment_groups(assembly, member_snapshot)
        composite_segments = classify_composite_main_material_segments(
            assembly,
            member_snapshot,
            main_material_groups=main_groups,
        )
        h_main_positions = {
            text(plate.part_position)
            for segment in composite_segments
            if segment.segment_type == CompositeSegmentType.H_OR_BH_SECTION
            for plate in segment.main_plates
            if text(plate.part_position)
        }

        if body_type == "H":
            for position in h_main_positions:
                _annotation(annotations, member, position)["main_material"] = "是"
            continue

        relations = classify_box_part_spatial_relations(assembly, member_snapshot, main_groups)
        part_position_by_id = {
            text(part.get("partId")): text(part.get("partPosition"))
            for part in assembly.get("parts", [])
            if text(part.get("partId")) and text(part.get("partPosition"))
        }
        for relation in relations:
            position = text(relation.part_position)
            if not position:
                continue
            if relation.relation_to_box_body == "MAIN_WALL":
                _annotation(annotations, member, position)["main_material"] = "是"
            elif body_type == "BOX" and relation.relation_to_box_body == "INSIDE_BODY":
                _annotation(annotations, member, position)["box_inside"] = "是"

        if body_type == "BOX":
            main_wall_ids = {
                text(part_id)
                for group in main_groups
                if group.group_type == "BOX_MAIN_WALL_CONFIRMED_SET"
                for part_id in group.part_ids
                if text(part_id)
            }
            outside_box_part_ids = {
                text(relation.part_id)
                for relation in relations
                if relation.relation_to_box_body == "OUTSIDE_ATTACHMENT" and text(relation.part_id)
            }
            appendages = classify_appendage_clusters_from_bundle(
                assembly,
                member_snapshot or {},
                body_part_ids=main_wall_ids,
                appendage_part_ids=outside_box_part_ids,
            )
            for appendage in appendages:
                label = _appendage_label(appendage.role, appendage.cluster_id)
                for part_id in appendage.part_ids:
                    position = part_position_by_id.get(text(part_id))
                    if position:
                        _annotation(annotations, member, position)["appendage"] = label

    return exported_members, member_body_types, annotations


def _annotation(
    annotations: dict[tuple[str, str], dict[str, Any]],
    member: str,
    position: str,
) -> dict[str, Any]:
    return annotations.setdefault(
        (member, position),
        {
            "main_material": "否",
            "box_inside": "否",
            "appendage": "",
        },
    )


def _appendage_label(role: str, cluster_id: str) -> str:
    if role == "Bracket":
        return f"牛腿({cluster_id})"
    if role == "ConnectionPlate":
        return f"连接板({cluster_id})"
    if role == "Unknown":
        return f"Unknown({cluster_id})"
    return f"{role}({cluster_id})" if role else ""


def model_delivery_role(annotation: dict[str, Any]) -> tuple[bool, str]:
    if annotation.get("main_material") == "是":
        return True, "主材"
    if annotation.get("box_inside") == "是":
        return True, "BOX内零件"
    return False, "后装配零件"


def check_workbook(excel_path: Path, cache_dir: Path, output_path: Path, summary_path: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(excel_path)
    values_wb = openpyxl.load_workbook(excel_path, data_only=True)
    member_ws = values_wb.worksheets[MEMBER_SHEET_INDEX]
    part_values_ws = values_wb.worksheets[PART_SHEET_INDEX]
    part_ws = wb.worksheets[PART_SHEET_INDEX]

    member_routes = load_member_routes(member_ws)
    exported_members, member_body_types, annotations = load_algorithm_annotations(cache_dir)

    first_added_col = part_ws.max_column + 1
    member_type_col = first_added_col
    main_material_col = first_added_col + 1
    box_inside_col = first_added_col + 2
    appendage_col = first_added_col + 3
    result_col = first_added_col + 4
    part_ws.cell(3, member_type_col).value = "构件类型判定"
    part_ws.cell(3, main_material_col).value = "主材判定"
    part_ws.cell(3, box_inside_col).value = "BOX内零件判定"
    part_ws.cell(3, appendage_col).value = "牛腿归属"
    part_ws.cell(3, result_col).value = "配送核对结果"

    counts: Counter[str] = Counter()
    checked_rows = 0
    skipped_not_exported = 0
    missing_route_rows = 0
    examples: list[dict[str, str]] = []

    for row in range(4, part_values_ws.max_row + 1):
        member = text(part_values_ws.cell(row, PART_MEMBER_COL).value)
        position = text(part_values_ws.cell(row, PART_POSITION_COL).value)
        if not member or not position:
            continue

        if exported_members and member not in exported_members:
            result = "未核对：本次Tekla导出未包含该构件"
            skipped_not_exported += 1
            part_ws.cell(row, member_type_col).value = ""
            part_ws.cell(row, main_material_col).value = ""
            part_ws.cell(row, box_inside_col).value = ""
            part_ws.cell(row, appendage_col).value = ""
            part_ws.cell(row, result_col).value = result
            counts[result] += 1
            continue

        route = member_routes.get(member)
        if not route:
            result = "未核对：构件清单缺少班组/流水线"
            missing_route_rows += 1
            part_ws.cell(row, member_type_col).value = member_body_types.get(member, "")
            part_ws.cell(row, main_material_col).value = ""
            part_ws.cell(row, box_inside_col).value = ""
            part_ws.cell(row, appendage_col).value = ""
            part_ws.cell(row, result_col).value = result
            counts[result] += 1
            continue

        actual = text(part_values_ws.cell(row, PART_DELIVERY_TARGET_COL).value)
        annotation = annotations.get(
            (member, position),
            {
                "main_material": "否",
                "box_inside": "否",
                "appendage": "",
            },
        )

        part_ws.cell(row, member_type_col).value = member_body_types.get(member, "")
        part_ws.cell(row, main_material_col).value = annotation.get("main_material", "否")
        part_ws.cell(row, box_inside_col).value = annotation.get("box_inside", "否")
        part_ws.cell(row, appendage_col).value = annotation.get("appendage", "")

        should_line, reason = model_delivery_role(annotation)
        expected = route["line"] if should_line else route["team"]
        if actual == expected:
            result = "一致"
        elif should_line:
            result = f"不一致：算法判为{reason}，应配送到{expected}，Excel为{actual or '空'}"
        else:
            result = f"不一致：算法判为后装配零件，应配送到{expected}，Excel为{actual or '空'}"

        part_ws.cell(row, result_col).value = result
        counts[result] += 1
        checked_rows += 1
        if result != "一致" and len(examples) < 20:
            examples.append(
                {
                    "构件": member,
                    "零件": position,
                    "原因": reason,
                    "应配送": expected,
                    "Excel": actual,
                    "结果": result,
                }
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    summary = {
        "excel": str(excel_path),
        "cache": str(cache_dir),
        "output_xlsx": str(output_path),
        "checked_part_rows": checked_rows,
        "skipped_not_exported_part_rows": skipped_not_exported,
        "missing_route_rows": missing_route_rows,
        "exported_member_count": len(exported_members),
        "member_route_count": len(member_routes),
        "result_counts": dict(counts),
        "examples": examples,
    }
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Check part delivery routing against member team/line rules.")
    parser.add_argument("--excel", type=Path, required=True, help="Input xlsx workbook.")
    parser.add_argument("--cache-dir", type=Path, required=True, help="Tekla export cache directory.")
    parser.add_argument("--output", type=Path, required=True, help="Output xlsx workbook.")
    parser.add_argument("--summary", type=Path, required=True, help="Output summary json.")
    args = parser.parse_args()

    summary = check_workbook(args.excel, args.cache_dir, args.output, args.summary)
    print(
        json.dumps(
            {
                "output_xlsx": summary["output_xlsx"],
                "summary": str(args.summary),
                "checked_part_rows": summary["checked_part_rows"],
                "skipped_not_exported_part_rows": summary["skipped_not_exported_part_rows"],
                "missing_route_rows": summary["missing_route_rows"],
                "result_counts": summary["result_counts"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
