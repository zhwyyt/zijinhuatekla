import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from zijinhua_tekla.reports.combined import (
    BUNDLE_FILENAME,
    export_combined_excel,
    inspect_export_root,
)


class CombinedReportTests(unittest.TestCase):
    def test_export_one_workbook_with_two_sheets(self):
        with TemporaryDirectory() as temp:
            root = Path(temp)
            assembly = {
                "assemblyId": "100",
                "mainPartId": "10",
                "metadata": {"assemblyPosition": "T3-MJ-1"},
                "parts": [
                    {
                        "partId": "10",
                        "partPosition": "T3-H-999",
                        "profileString": "BH300*150*6*8",
                        "runtimeType": "Beam",
                        "isPlateLike": False,
                        "boltHoleCount": 1,
                        "booleanCutCount": 1,
                        "booleanCutDetails": [
                            {
                                "operativePartId": 20,
                                "cutsFatherSolid": True,
                                "cutProof": "GET_CUT_PART",
                            }
                        ],
                    }
                ],
                "relationships": [],
            }
            bundle = {"assemblies": [assembly]}
            (root / BUNDLE_FILENAME).write_text(
                json.dumps(bundle, ensure_ascii=False),
                encoding="utf-8",
            )
            members = root / "members"
            members.mkdir()
            (members / "member_T3-MJ-1.json").write_text(
                json.dumps({"Classification": {"MainClass": 1}}, ensure_ascii=False),
                encoding="utf-8",
            )
            output = root / "out"

            info = inspect_export_root(root)
            result = export_combined_excel(root, "项目:A", "区域/B", output)
            workbook = load_workbook(result.path, read_only=True)

            self.assertTrue(info["valid"])
            self.assertEqual(1, result.assembly_count)
            self.assertEqual(1, result.feature_row_count)
            self.assertEqual(1, result.complexity_row_count)
            self.assertEqual(0, result.error_count)
            self.assertEqual("项目_A_区域_B.xlsx", result.path.name)
            self.assertEqual(["零件特征", "构件复杂度"], workbook.sheetnames)
            workbook.close()

    def test_invalid_root_is_rejected(self):
        with TemporaryDirectory() as temp:
            with self.assertRaises(FileNotFoundError):
                export_combined_excel(temp, "项目", "区域")




if __name__ == "__main__":
    unittest.main()

