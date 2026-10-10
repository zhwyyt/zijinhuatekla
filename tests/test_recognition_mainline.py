import copy
from contextlib import closing
from io import BytesIO
import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import mock

from openpyxl import load_workbook

from zijinhua_tekla.classifiers.member_structural_role import classify_body_bracket_assembly
from zijinhua_tekla.reports.combined import BUNDLE_FILENAME, export_combined_excel
from zijinhua_tekla.reports.member_complexity import _member_axis
from zijinhua_tekla.reports.member_role import ALIGNED_SPAN_AXIS_SOURCE, member_role_row
from zijinhua_tekla.pipeline.offline import run_model_recognition
from zijinhua_tekla.reports.recognition import write_recognition_workbook


def column_fixture():
    assembly = {
        "assemblyId": "A1",
        "mainPartId": 1,
        "metadata": {
            "assemblyPosition": "SAMPLE",
            "memberAxisEvidence": {
                "origin": {"x": 0, "y": 0, "z": 0},
                "direction": {"x": 0, "y": 0, "z": 1},
                "length": 4000,
                "source": ALIGNED_SPAN_AXIS_SOURCE,
            },
            "externalConnectionEvidence": {
                "links": [{"axisLocation": "end", "otherAxisPose": "horizontal", "relation": "Bolt"}],
            },
        },
        "parts": [
            {
                "partId": 1, "partPosition": "P1", "profileString": "BH400*200*8*12",
                "runtimeType": "Beam", "isPlateLike": False,
                "obbDims": {"x": 200, "y": 400, "z": 4000},
            },
        ],
        "relationships": [],
    }
    member = {
        "Member": {"Position": "SAMPLE", "AssemblyId": "A1", "MainAxis": {"Direction": {"X": 0, "Y": 0, "Z": 1}}},
        "LongitudinalAxisSource": ALIGNED_SPAN_AXIS_SOURCE,
        "AxisSegments": [{"Length": 4000, "Direction": {"X": 0, "Y": 0, "Z": 1}}],
        "GuidePolyline": [{"X": 0, "Y": 0, "Z": 0}, {"X": 0, "Y": 0, "Z": 4000}],
    }
    return assembly, member


def write_export(root, pairs):
    (root / BUNDLE_FILENAME).write_text(json.dumps({"assemblies": [a for a, _ in pairs]}), encoding="utf-8")
    (root / "members").mkdir()
    for assembly, member in pairs:
        name = assembly["metadata"]["assemblyPosition"]
        (root / "members" / f"member_{name}.json").write_text(json.dumps(member), encoding="utf-8")


class RecognitionMainlineTests(unittest.TestCase):
    def _same_position_instances(self, root, with_batch=True):
        vertical, first = column_fixture()
        horizontal, second = copy.deepcopy((vertical, first))
        horizontal["assemblyId"] = "A2"
        horizontal["mainPartId"] = 2
        horizontal["parts"][0]["partId"] = 2
        horizontal["metadata"]["memberAxisEvidence"]["direction"] = {"x": 1, "y": 0, "z": 0}
        second["Member"]["AssemblyId"] = "A2"
        second["Member"]["MainAxis"]["Direction"] = {"X": 1, "Y": 0, "Z": 0}
        second["AxisSegments"][0]["Direction"] = {"X": 1, "Y": 0, "Z": 0}
        second["GuidePolyline"] = [{"X": 0, "Y": 0, "Z": 0}, {"X": 4000, "Y": 0, "Z": 0}]
        write_export(root, [(vertical, first), (horizontal, second)])
        if with_batch:
            (root / "members/batch_summary.json").write_text(
                json.dumps({"Members": [first, second]}), encoding="utf-8",
            )

    def test_same_position_instances_keep_their_own_axis_in_batch_report(self):
        with TemporaryDirectory() as temp:
            root = Path(temp)
            self._same_position_instances(root)
            result = export_combined_excel(root, "项目", "同号", root / "out")
            self.assertEqual((2, 2, 2, 0), (
                result.role_row_count, result.complexity_row_count,
                result.feature_row_count, result.error_count,
            ))
            with closing(load_workbook(BytesIO(result.path.read_bytes()))) as workbook:
                rows = iter(workbook["构件角色"].values)
                headers = next(rows)
                roles = {row["装配ID"]: row for row in (dict(zip(headers, values)) for values in rows)}
                self.assertEqual("竖直", roles["A1"]["轴线姿态"])
                self.assertEqual("水平", roles["A2"]["轴线姿态"])

    def test_overwritten_instance_without_batch_is_a_failure_not_a_substitution(self):
        with TemporaryDirectory() as temp:
            root = Path(temp)
            self._same_position_instances(root, with_batch=False)
            result = export_combined_excel(root, "项目", "身份缺失", root / "out")
            self.assertEqual((1, 1, 1, 1), (
                result.role_row_count, result.complexity_row_count,
                result.feature_row_count, result.error_count,
            ))
            self.assertEqual("A1", result.errors[0]["assembly_id"])
            self.assertIn("装配ID不一致", result.errors[0]["error"])

    def test_single_member_explicit_assembly_reads_its_own_batch_snapshot(self):
        with TemporaryDirectory() as temp:
            root = Path(temp)
            self._same_position_instances(root)
            result = run_model_recognition(root, "SAMPLE", "A1")
            self.assertEqual("A1", result.member["Member"]["AssemblyId"])
            self.assertEqual("vertical", result.member_structural_role.pose)

    def test_accepted_column_role_is_visible_in_formal_workbook(self):
        assembly, member = column_fixture()
        with TemporaryDirectory() as temp:
            root = Path(temp)
            write_export(root, [(assembly, member)])
            result = export_combined_excel(root, "项目", "区域", root / "out")
            with closing(load_workbook(BytesIO(result.path.read_bytes()))) as workbook:
                self.assertEqual(["构件角色", "构件复杂度", "零件特征"], workbook.sheetnames)
                headers, values = list(workbook["构件角色"].values)
                role = dict(zip(headers, values))
                self.assertEqual("COLUMN", role["角色代码"])
                self.assertEqual("框架柱", role["职能子类"])
                self.assertEqual("MR-COL-003", role["规则"])
                self.assertEqual(ALIGNED_SPAN_AXIS_SOURCE, role["主轴来源"])
            self.assertEqual((1, 1, 1, 0, 0), (
                result.role_row_count, result.complexity_row_count, result.feature_row_count,
                result.error_count, result.axis_review_count,
            ))

    def test_real_polyline_is_not_replaced_by_the_bundle_main_segment(self):
        assembly, member = column_fixture()
        member["AxisSegments"] = [
            {"Length": 2236, "Direction": {"X": 1, "Y": 0, "Z": 2}},
            {"Length": 2236, "Direction": {"X": -1, "Y": 0, "Z": 2}},
        ]
        member["GuidePolyline"].insert(1, {"X": 1000, "Y": 0, "Z": 2000})
        role = classify_body_bracket_assembly(assembly, member)
        self.assertEqual(("COLUMN", "vertical", "折线", 1), (role.role, role.pose, role.axis_form, role.bend_count))
        self.assertEqual((0, 0, 1), _member_axis(member))
        with TemporaryDirectory() as temp:
            root = Path(temp)
            write_export(root, [(assembly, member)])
            result = export_combined_excel(root, "项目", "折线", root / "out")
            with closing(load_workbook(BytesIO(result.path.read_bytes()))) as workbook:
                headers, values = list(workbook["构件复杂度"].values)
                complexity = dict(zip(headers, values))
                self.assertEqual(("折线", 1), (complexity["主材轴线形状"], complexity["折弯数量"]))

    def test_old_axis_and_mixed_sources_require_reexport(self):
        assembly, member = column_fixture()
        assembly["metadata"]["memberAxisEvidence"]["source"] = "mainPart.longestAxis"
        self.assertIn("重新导出", member_role_row(assembly, member)["复核状态"])
        assembly["metadata"]["memberAxisEvidence"]["source"] = ALIGNED_SPAN_AXIS_SOURCE
        member["LongitudinalAxisSource"] = "legacy"
        self.assertIn("重新导出", member_role_row(assembly, member)["复核状态"])

    def test_single_member_pipeline_and_json_keep_the_accepted_role(self):
        assembly, member = column_fixture()
        with TemporaryDirectory() as temp:
            root = Path(temp)
            write_export(root, [(assembly, member)])
            result = run_model_recognition(root, "SAMPLE")
            self.assertEqual("COLUMN", result.member_structural_role.role)
            report = write_recognition_workbook(result, root / "out", "SAMPLE")
            data = json.loads(report.json_path.read_text(encoding="utf-8"))
            self.assertEqual("COLUMN", data["member_structural_role"]["role"])
            self.assertEqual(ALIGNED_SPAN_AXIS_SOURCE, data["member_axis_source"])
            with closing(load_workbook(BytesIO(report.xlsx_path.read_bytes()))) as workbook:
                self.assertIn("构件角色", workbook.sheetnames)

    def test_missing_connection_is_reported_without_changing_the_node_role(self):
        from tests.test_connection_review import _brace, _node

        brace, node = _brace(), _node((0, 180, 0))
        pairs = []
        for assembly in (brace, node):
            axis = assembly["metadata"]["memberAxisEvidence"]
            axis["source"] = ALIGNED_SPAN_AXIS_SOURCE
            origin, direction, length = axis["origin"], axis["direction"], axis["length"]
            member = {
                "LongitudinalAxisSource": ALIGNED_SPAN_AXIS_SOURCE,
                "AxisSegments": [{"Length": length, "Direction": {k.upper(): v for k, v in direction.items()}}],
                "GuidePolyline": [
                    {k.upper(): v for k, v in origin.items()},
                    {k.upper(): origin[k] + direction[k] * length for k in origin},
                ],
            }
            pairs.append((assembly, member))
        with TemporaryDirectory() as temp:
            root = Path(temp)
            write_export(root, pairs)
            result = export_combined_excel(root, "项目", "连接复核", root / "out")
            self.assertEqual(1, result.connection_review_count)
            with closing(load_workbook(BytesIO(result.path.read_bytes()))) as workbook:
                self.assertIn("连接复核", workbook.sheetnames)
                rows = list(workbook["构件角色"].values)
                roles = [dict(zip(rows[0], row))["角色代码"] for row in rows[1:]]
                self.assertEqual(["BRACE", "UNKNOWN"], roles)
    def test_member_mark_never_replaces_connection_evidence(self):
        assembly, member = column_fixture()
        assembly["metadata"]["assemblyPosition"] = "GKZ-GL-MJ"
        assembly["metadata"]["externalConnectionEvidence"]["links"] = []
        role = member_role_row(assembly, member)
        self.assertEqual("UNKNOWN", role["角色代码"])
        self.assertNotIn("GKZ-GL-MJ", role["证据"])

    def test_failed_assembly_does_not_leave_partial_success_rows(self):
        good, member = column_fixture()
        bad = copy.deepcopy(good)
        bad["assemblyId"] = "A2"
        bad["metadata"]["assemblyPosition"] = "BAD"
        from zijinhua_tekla.reports.combined import classify_member_complexity

        def classify(assembly, document):
            if assembly["assemblyId"] == "A2":
                raise ValueError("失败样例")
            return classify_member_complexity(assembly, document)

        with TemporaryDirectory() as temp:
            root = Path(temp)
            write_export(root, [(good, member), (bad, member)])
            with mock.patch("zijinhua_tekla.reports.combined.classify_member_complexity", side_effect=classify):
                result = export_combined_excel(root, "项目", "失败", root / "out")
            self.assertEqual((1, 1, 1, 1), (
                result.role_row_count, result.complexity_row_count, result.feature_row_count, result.error_count,
            ))
            with closing(load_workbook(BytesIO(result.path.read_bytes()))) as workbook:
                self.assertIn("识别失败", workbook.sheetnames)
                self.assertEqual("BAD", workbook["识别失败"]["A2"].value)


if __name__ == "__main__":
    unittest.main()
